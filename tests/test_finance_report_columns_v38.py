"""財務欄位及下載驗證；全部為虛構資料，不連正式資料庫。"""
import ast
from copy import deepcopy
from datetime import date
from pathlib import Path
import unittest
from io import BytesIO
from openpyxl import load_workbook
from pandas.testing import assert_frame_equal
from test_financial_backup import load_functions
from monthly_report_layout import BONUS_COLUMNS

import test_online_finance as base


class FinanceReportColumnsTests(unittest.TestCase):
    app=base.FinancialScreenTests.app

    def monthly(self,data):
        app=self.app('financial',tables=data,financial_report_view='每月報表')
        app.date_input[0].set_value(date(2026,9,1))
        app.date_input[1].set_value(date(2026,9,30))
        app.button[0].click().run()
        self.assertFalse(app.exception)
        return app

    def test_usage_makeup_after_hours_same_in_export(self):
        data=base.tables();original=data['session_usages'][0]
        data['session_usages']=[dict(original,id='u'+str(i),session_seq=i+1,is_makeup=flag)
                                for i,flag in enumerate((True,False,None))]
        app=self.app('general',tables=data,general_query_main_tab='銷課查詢',general_usage_mode='日期區間')
        app.date_input[0].set_value(date(2026,9,1));app.date_input[1].set_value(date(2026,9,30))
        app.button[0].click().run()
        self.assertFalse(app.exception)
        frame=app.dataframe[0].value
        self.assertEqual(frame.columns.tolist()[-2:],['時數','補單'])
        self.assertEqual(frame['補單'].tolist(),['是','否','否'])
        exported=app.session_state['exports']['銷課查詢']
        self.assertEqual(exported.to_dict('records'),frame.to_dict('records'))
        self.assertEqual(frame['銷課金額（未稅）'].sum(),2001)

    def test_usage_query_fetches_persisted_flag(self):
        tree=ast.parse((Path(__file__).resolve().parents[1]/'app.py').read_text(encoding='utf-8'))
        page=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='general_queries_page')
        renderer=next(n for n in page.body if isinstance(n,ast.FunctionDef) and n.name=='render_usage_results')
        selections=[n.args[0].value for n in ast.walk(renderer) if isinstance(n,ast.Call)
                    and isinstance(n.func,ast.Attribute) and n.func.attr=='select'
                    and n.args and isinstance(n.args[0],ast.Constant)]
        self.assertTrue(any('is_makeup' in text.split(',') for text in selections))

    def test_monthly_accounts_first_and_duplicate_names_stay_separate(self):
        data=base.tables();data['profiles'][0]['username']='SMC001'
        data['profiles'].append(dict(data['profiles'][0],id='c2',username='SMC002'))
        data['purchases'].append(dict(data['purchases'][0],id='p2',coach_id='c2',purchase_code='20260901-002',session_hours=1.5))
        data['session_usages'].append(dict(data['session_usages'][0],id='u2',purchase_id='p2',coach_id='c2',deducted_amount=350,deducted_net_amount=333))
        app=self.monthly(data)
        for key in ('每月教練時數','每月教練營收'):
            frame=app.session_state['exports'][key]
            self.assertEqual(frame.columns.tolist()[:2],['帳號','教練'])
            self.assertEqual(set(frame['帳號']),{'SMC001','SMC002'})
            self.assertEqual(len(frame),2)
        hours=app.session_state['exports']['每月教練時數'].set_index('帳號')
        self.assertEqual(hours.loc['SMC001','可計執行時數'],3)
        self.assertEqual(hours.loc['SMC002','可計執行時數'],1.5)
        revenue=app.session_state['exports']['每月教練營收'].set_index('帳號')
        self.assertEqual(revenue.loc['SMC001','金額總計（未稅）'],867)
        self.assertEqual(revenue.loc['SMC002','金額總計（未稅）'],333)

    def test_missing_username_is_not_invented(self):
        app=self.monthly(base.tables())
        self.assertEqual(app.session_state['exports']['每月教練時數'].iloc[0]['帳號'],'')

    def test_cross_month_hours_separate_without_changing_execution_or_revenue(self):
        data=base.tables();data['course_terminations']=[]
        original=data['session_usages'][0]
        data['session_usages']=[dict(original,is_makeup=False),
            dict(original,id='u2',usage_date='2026-09-03',actual_usage_date='2026-09-01',is_makeup=True,session_seq=2),
            dict(original,id='u3',usage_date='2026-09-05',actual_usage_date='2026-08-31',is_makeup=True,session_seq=3,deducted_net_amount=666)]
        data['purchases'].append(dict(data['purchases'][0],id='p2',purchase_code='20260901-002',course_name='動磁波',session_hours=0.5,total_amount=1050,total_sessions=2))
        data['session_usages'].append(dict(original,id='m1',purchase_id='p2',usage_date='2026-09-06',actual_usage_date='2026-08-30',is_makeup=True,deducted_amount=525,deducted_net_amount=500))
        app=self.monthly(data);frames=app.session_state['exports'];hours=frames['每月教練時數']
        columns=hours.columns.tolist()
        self.assertEqual(columns[columns.index('可計執行時數')+1],'補單')
        self.assertEqual(hours.iloc[0]['補單'],1.5)
        self.assertEqual(hours.iloc[0]['可計執行時數'],4)
        self.assertEqual(frames['每月教練營收'].iloc[0]['銷課（未稅）'],2500)
        backup=load_functions()['_financial_backup_frames'](data)
        self.assertEqual(backup['財務-教練時數'].iloc[0]['補單'],1.5)
        self.assertEqual(backup['財務-教練時數'].iloc[0]['可計執行時數'],4)

    def bonus_data(self):
        data=base.tables();data['profiles'][0]['username']='SMC001'
        data['profiles'].append(dict(data['profiles'][0],id='c2',username='SMC002'))
        original=data['purchases'][0]
        data['purchases']=[dict(original,total_sessions=1,status='completed'),
            dict(original,id='p2',purchase_code='20260912-001',purchase_date='2026-09-12',total_amount=1050,total_sessions=1,referral='虛構轉介'),
            dict(original,id='p3',purchase_code='20260915-001',purchase_date='2026-09-15',coach_id='c2',total_amount=3150,total_sessions=1)]
        data['session_usages']=[dict(data['session_usages'][0],deducted_amount=2100,deducted_net_amount=2000),
            dict(data['session_usages'][0],id='u2',purchase_id='p2',usage_date='2026-09-13',actual_usage_date='2026-09-13',deducted_amount=1050,deducted_net_amount=1000),
            dict(data['session_usages'][0],id='u3',purchase_id='p3',coach_id='c2',usage_date='2026-09-15',actual_usage_date='2026-09-15',deducted_amount=3150,deducted_net_amount=3000)]
        data['course_terminations']=[]
        data['bonus_rules']=[{'id':'r','rule_name':'虛構規則','active':True,'effective_from':'2026-01-01',
            'talk_rate':0.1,'completion_rate':0.05,'referral_first_talk_eligible':False,'referral_first_completion_eligible':False}]
        # 完整備份須有各筆真實的虛構收款來源，不能借用不符合約的原付款。
        data['purchase_payments']=[dict(data['purchase_payments'][0],id='pay'+str(i),purchase_id=p['id'],amount=p['total_amount'],
            net_amount=n,paid_date=p['purchase_date']) for i,(p,n) in enumerate(zip(data['purchases'],(2000,1000,3000)))]
        return data

    def test_bonus_five_columns_ids_reconciliation_and_backup_match(self):
        data=self.bonus_data();app=self.monthly(data);frames=app.session_state['exports']
        backup=load_functions()['_financial_backup_frames'](data)
        for sheet,total in [('談單獎金總計',200),('結單獎金總計',100)]:
            frame=frames[sheet]
            self.assertEqual(frame.columns.tolist(),BONUS_COLUMNS)
            by_account=frame.set_index('帳號')
            self.assertEqual(by_account.loc['SMC001','不計成交金額'],1000)
            self.assertEqual(by_account.loc['SMC001','可計成交金額'],2000)
            self.assertEqual(by_account.loc['SMC001','獎金合計'],total)
            self.assertEqual(by_account.loc['SMC002','可計成交金額'],3000)
            self.assertEqual(by_account.loc['SMC002','獎金合計'],total*1.5)
            assert_frame_equal(frame,backup['財務-'+sheet])
        book=load_workbook(BytesIO(app.session_state['export_bytes']),data_only=True)
        self.assertEqual([c.value for c in book['談單獎金總計'][1]],BONUS_COLUMNS)
        for key in ('每月教練談單獎金','每月教練結單獎金'):
            self.assertFalse(any(c.startswith('_') for c in frames[key].columns))
            self.assertFalse(any(c.startswith('_') for c in backup['財務-教練'+('談單獎金' if '談單' in key else '結單獎金')].columns))

    def test_missing_rule_is_pending_not_ineligible(self):
        data=self.bonus_data();data['bonus_rules'][0]['effective_from']='2026-09-10'
        app=self.monthly(data);frames=app.session_state['exports']
        for summary,pending in [('談單獎金總計','談單待確認'),('結單獎金總計','結單待確認')]:
            row=frames[summary].set_index('帳號').loc['SMC001']
            self.assertEqual(row['不計成交金額'],1000)
            self.assertEqual(row['可計成交金額'],0)
            self.assertEqual(row['獎金合計'],0)
            self.assertEqual(frames[pending]['待確認成交金額（未稅）'].sum(),2000)
            self.assertEqual(len(frames[pending]),1)
        self.assertTrue(any('另列待確認' in item.value for item in app.warning))

    def test_zero_bonus_rate_is_still_eligible_amount(self):
        data=self.bonus_data();data['bonus_rules'][0].update(talk_rate=0,completion_rate=0)
        app=self.monthly(data)
        for key in ('談單獎金總計','結單獎金總計'):
            row=app.session_state['exports'][key].set_index('帳號').loc['SMC001']
            self.assertEqual(row['可計成交金額'],2000)
            self.assertEqual(row['不計成交金額'],1000)
            self.assertEqual(row['獎金合計'],0)


if __name__=='__main__':unittest.main()
