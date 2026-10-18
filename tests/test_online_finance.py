"""正式規格的報表與下載回歸，全部使用虛構資料，不呼叫正式 DB。"""
import ast
from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import unittest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from finance_reporting import display_money, payment_allocations, usage_net, termination_money, ledger_balance, event_net_total
from project_queries import execution_report, project_report_frame
from test_confirmed_queries import SCRIPT

ROOT=Path(__file__).resolve().parents[1]


def tables():
    purchase={'id':'p1','purchase_code':'20260901-001','member_id':'m1','coach_id':'c1',
              'course_name':'虛構課程','total_amount':2100,'total_sessions':3,'session_hours':1,
              'purchase_date':'2026-09-01','expiry_date':'2027-09-01','status':'cancelled',
              'purchase_kind':'first','payment_plan':'full','installment_count':1}
    term={'id':'t1','purchase_id':'p1','termination_date':'2026-09-30','termination_type':'refund',
          'remaining_sessions':2,'remaining_amount':1400,'fee_amount':280,'refund_amount':1120,
          'recognized_amount':280,'remaining_net_amount':1333,'fee_net_amount':267,
          'refund_net_amount':1066,'recognized_net_amount':267,'completion_bonus_eligible':False}
    return {'profiles':[{'id':'c1','display_name':'虛構教練','role':'coach','active':True}],
            'members':[{'id':'m1','member_name':'虛構會員'}],'purchases':[purchase],
            'purchase_payments':[{'id':'pay1','purchase_id':'p1','amount':2100,'net_amount':2000,
                                  'paid_date':'2026-09-01','installment_no':1}],
            'session_usages':[{'id':'u1','purchase_id':'p1','coach_id':'c1','session_seq':1,
                               'usage_date':'2026-09-02','actual_usage_date':'2026-09-02',
                               'deducted_amount':700,'deducted_net_amount':667}],
            'course_terminations':[term],
            'trial_items':[{'id':'tr1','coach_id':'c1','entry_date':'2026-09-03','amount':105,
                            'hours':1,'member_name':'虛構會員','course_type':'訓練','content':'虛構體驗'}],
            'single_sales':[{'id':'s1','coach_id':'c1','entry_date':'2026-09-04','amount':105,
                            'hours':1,'member_name':'虛構會員','course_type':'訓練','content':'虛構單次'}]}


class PureReportingTests(unittest.TestCase):
    def test_rounding_per_event_not_total(self):
        self.assertEqual(event_net_total([{'amount':10.01}]*3,'amount'),30)
        self.assertEqual(display_money(30.03,'未稅'),29)

    def test_saved_usage_and_missing_is_error(self):
        self.assertEqual(usage_net({'deducted_net_amount':666}),666)
        with self.assertRaises(ValueError): usage_net({'deducted_amount':700})

    def test_term_uses_saved_split_and_balance_not_gross_redivision(self):
        term=tables()['course_terminations'][0]
        self.assertEqual(termination_money(term,'refund_amount','未稅'),1066)
        self.assertEqual(display_money(term['refund_amount'],'未稅'),1067)
        self.assertEqual(ledger_balance(2000,667,term),0)
        self.assertEqual(ledger_balance(1999,2000),-1)

    def test_payment_saved_value_checked_and_tail(self):
        p=[{'id':'p','total_amount':2100}]
        rows=[{'id':str(i),'purchase_id':'p','amount':700,'paid_date':f'2026-09-0{i+1}'} for i in range(3)]
        self.assertEqual([x['_net_amount'] for x in payment_allocations(p,rows)],[667,667,666])
        rows[0]['net_amount']=665
        with self.assertRaises(ValueError): payment_allocations(p,rows)

    def test_project_balance_is_net_event_ledger(self):
        project={'id':'p','project_name':'虛構專案','funding_type':'stored'}
        deposits=[{'id':'d','project_id':'p','deposit_date':'2026-09-01','amount':210,'transaction_type':'opening'}]
        entries=[{'id':str(i),'project_id':'p','entry_date':'2026-09-02','line_amount':'10.01',
                  'item_hours':1,'quantity':1} for i in range(3)]
        result=execution_report(project,entries,deposits,{},display_money,mode='cutoff',end='2026-09-30')
        self.assertEqual(result['balance_net'],170)
        self.assertEqual(display_money(result['balance_gross'],'未稅'),171)
        self.assertEqual(project_report_frame(result,'cutoff').iloc[-1]['金額（未稅）'],170)

    def test_no_parallel_tax_formula_in_app(self):
        text=(ROOT/'app.py').read_text(encoding='utf-8')
        tree=ast.parse(text)
        fn=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=='_tax_display_amount')
        self.assertIn('display_money',ast.unparse(fn))
        self.assertNotIn('1.05',ast.unparse(fn))
        self.assertNotIn('prepaid_net_balance=max(',text)
        self.assertNotIn('project_prepaid_total=_tax_display_amount',text)


class FinancialScreenTests(unittest.TestCase):
    def app(self,view,**state):
        script=SCRIPT.replace('me={"id"',
            'original_writer=_excel_bytes\ndef _excel_bytes(*a,**k):\n value=original_writer(*a,**k)\n st.session_state["exports"]=a[0]\n st.session_state["export_bytes"]=value\n return value\nme={"id"',1)
        app=AppTest.from_string('APP_PATH='+repr(str(ROOT/'app.py'))+'\n'+script,default_timeout=60)
        for key,value in {'view':view,'role':'admin','tables':tables(),**state}.items():
            app.session_state[key]=value
        app.run();self.assertFalse(app.exception)
        return app

    def test_general_balance_saved_refund_detail_and_excel(self):
        app=self.app('general',general_query_main_tab='預收餘額查詢',general_balance_mode='預收餘額（未稅）')
        # 使用實際元件，不依內部下載字串假設欄位。
        app.date_input[0].set_value(date(2026,9,30));app.button[0].click().run()
        self.assertFalse(app.exception)
        frame=app.dataframe[0].value
        self.assertEqual(frame.iloc[0]['餘額（未稅）'],0)
        self.assertEqual(frame.iloc[0]['退費金額（未稅）'],1066)
        book=load_workbook(BytesIO(app.session_state['export_bytes']),data_only=True)
        headers=[c.value for c in book.worksheets[0][1]]
        self.assertEqual(book.worksheets[0].cell(2,headers.index('餘額（未稅）')+1).value,0)

    def test_member_report_retains_negative_and_matches_general(self):
        app=self.app('financial',financial_report_view='會員報表')
        app.date_input[0].set_value(date(2026,9,1));app.date_input[1].set_value(date(2026,9,30))
        app.button[0].click().run();self.assertFalse(app.exception)
        sheets=app.session_state['exports']
        frame=sheets['預收餘額明細']
        self.assertEqual(frame.iloc[0]['實際預收剩餘金額'],0)
        self.assertEqual(frame.iloc[0]['累計銷課金額'],667)
        self.assertEqual(sheets['成交預收彙總'].iloc[0]['成交總金額'],2000)

    def test_monthly_net_termination_and_operation_revenue(self):
        app=self.app('financial',financial_report_view='每月報表')
        app.date_input[0].set_value(date(2026,9,1));app.date_input[1].set_value(date(2026,9,30))
        app.button[0].click().run();self.assertFalse(app.exception)
        sheets=app.session_state['exports']
        self.assertEqual(sheets['每月課程中止'].iloc[0]['實際退費金額（未稅）'],1066)
        self.assertEqual(sheets['每月銷課']['銷課金額'].sum(),667)
        row=sheets['每月教練營收'].iloc[0]
        self.assertEqual(row['體驗項目金額（未稅）'],100)
        self.assertEqual(row['單堂銷售金額（未稅）'],100)

    def test_exact_runtime_login_imports_without_missing_module(self):
        app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=60).run()
        self.assertFalse(app.exception)
        self.assertGreaterEqual(len(app.text_input),2)
