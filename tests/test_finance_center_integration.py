"""接入同一份計算結果的畫面／下載、查詢與角色隔離。"""
from datetime import date
from io import BytesIO
from pathlib import Path
import sys
import unittest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tests'))
import test_query_views_v1_13_14 as previous_tests
from monthly_operation_queries import operation_report_frame,NET_COLUMN,MONEY_COLUMN,BASIS_COLUMN
from financial_calculations import build_operation_result


class FinanceCenterUI(previous_tests.OperationQueryUITests):
    def export_app(self,records):
        marker='me={"id"'
        self.assertIn(marker,previous_tests.OP_SCRIPT)
        script=previous_tests.OP_SCRIPT.replace(marker,
            'original_excel_writer=_excel_bytes\ndef _excel_bytes(*args,**kwargs):\n payload=original_excel_writer(*args,**kwargs)\n st.session_state["finance_export"]=payload\n return payload\n'+marker,1)
        app=AppTest.from_string('APP_PATH='+repr(str(ROOT/'app.py'))+'\n'+script,default_timeout=30)
        for key,value in {'view':'financial','role':'admin','financial_report_view':'體驗項目',
                          'monthly_operation_trial_items_mode':'截止日期','tables':{'trial_items':records}}.items():
            app.session_state[key]=value
        app.run();self.assertFalse(app.exception)
        app.date_input[0].set_value(date(2026,9,30));app.button[0].click().run()
        self.assertFalse(app.exception)
        return app

    def test_actual_download_summary_matches_screen(self):
        app=self.export_app([{'id':str(i),'entry_date':'2026-09-01','amount':'10.01'} for i in range(3)])
        book=load_workbook(BytesIO(app.session_state['finance_export']),data_only=True)
        summary=book['查詢摘要'];columns=[c.value for c in summary[1]]
        self.assertEqual(summary.cell(2,columns.index('體驗總計（未稅）')+1).value,30)
        self.assertEqual(summary.cell(2,columns.index('體驗總計（含稅）')+1).value,30.03)
        detail=book['體驗項目'];headers=[c.value for c in detail[1]]
        self.assertEqual(sum(detail.cell(r,headers.index(NET_COLUMN)+1).value for r in range(2,5)),30)
        self.assertEqual(app.metric[1].value,'$ 30')
        self.assertEqual(summary.cell(2,columns.index('未稅待確認筆數')+1).value,0)

    def test_actual_download_applies_same_standard_to_prior_history(self):
        app=self.export_app([{'id':'known','entry_date':'2026-09-01','amount':105},
                             {'id':'unknown','entry_date':'2026-08-31','amount':105}])
        book=load_workbook(BytesIO(app.session_state['finance_export']),data_only=True)
        summary=book['查詢摘要'];columns=[c.value for c in summary[1]]
        self.assertEqual(summary.cell(2,columns.index('體驗總計（未稅）')+1).value,200)
        self.assertEqual(summary.cell(2,columns.index('未稅待確認筆數')+1).value,0)
        self.assertEqual(app.metric[1].value,'$ 200')

    def test_confirmed_month_total_frame_and_export_are_same(self):
        records=[{'id':'a','entry_date':'2026-09-01','amount':'10.01','coach_id':'c1'},
                 {'id':'b','entry_date':'2026-09-30','amount':'10.01','coach_id':'c1'}]
        for report in ['體驗項目','單次銷售']:
            with self.subTest(report=report):
                app=self.app(report,records=records)
                app.date_input[0].set_value(date(2026,9,1));app.date_input[1].set_value(date(2026,9,30))
                app.button[0].click().run()
                self.assertFalse(app.exception)
                self.assertEqual(app.metric[0].value,'$ 20.02')
                self.assertEqual(app.metric[1].value,'$ 20')
                frame=app.dataframe[0].value
                self.assertEqual(frame[NET_COLUMN].tolist(),[10,10])
                self.assertTrue(all('重建' in x for x in frame[BASIS_COLUMN]))
                data=previous_tests.load_excel_writer()({'結果':frame})
                ws=load_workbook(BytesIO(data),data_only=True)['結果']
                heads=[c.value for c in ws[1]]
                net_column=heads.index(NET_COLUMN)+1
                self.assertEqual(sum(ws.cell(r,net_column).value for r in (2,3)),20)

    def test_mixed_cutoff_all_history_uses_unified_standard(self):
        app=self.app('體驗項目','截止日期',records=self.records())
        app.date_input[0].set_value(date(2026,9,30));app.button[0].click().run()
        self.assertFalse(app.exception)
        self.assertEqual(app.metric[1].value,'$ 3,095')
        self.assertFalse(any('不提供部分合計' in w.value for w in app.warning))

    def test_missing_amount_no_totals_or_download(self):
        for bad in [None,'NaN','']:
            with self.subTest(value=bad):
                app=self.app('體驗項目','截止日期',records=[{'id':'a','entry_date':'2026-09-01','amount':bad}])
                app.date_input[0].set_value(date(2026,9,30));app.button[0].click().run()
                self.assertFalse(app.exception)
                self.assertEqual(len(app.error),1)
                self.assertEqual(len(app.metric),0)
                self.assertEqual(len(app.get('download_button')),0)

    def test_duplicate_records_block_download(self):
        row={'id':'same','entry_date':'2026-09-01','amount':105}
        app=self.app('體驗項目','截止日期',records=[row,dict(row)])
        app.date_input[0].set_value(date(2026,9,30));app.button[0].click().run()
        self.assertEqual(len(app.error),1)
        self.assertEqual(len(app.get('download_button')),0)

    def test_negative_amount_remains_visible(self):
        app=self.app('體驗項目','截止日期',records=[{'id':'a','entry_date':'2026-09-01','amount':-105}])
        app.date_input[0].set_value(date(2026,9,30));app.button[0].click().run()
        self.assertEqual(app.metric[1].value,'$ -100')
        self.assertTrue(any('負金額' in w.value for w in app.warning))


if __name__=='__main__':unittest.main()
