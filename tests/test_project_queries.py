"""專案報表只使用虛構資料，核對日期、類型、餘額、時數與 Excel。"""
import ast
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
import sys
import unittest

import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from project_queries import execution_report, deposit_report, project_excel_number_formats, project_report_frame


def tax(value, mode):
    return int((Decimal(str(value)) / Decimal('1.05')).quantize(Decimal('1'), rounding=ROUND_HALF_UP))


def fixtures():
    project = {'id': 'p', 'project_name': '虛構儲值專案', 'funding_type': 'stored', 'stored_amount': 999999, 'stored_date': '2026-07-01'}
    deposits = [
        {'id': 'd1', 'project_id': 'p', 'deposit_date': '2026-07-01', 'amount': 10500, 'transaction_type': 'opening'},
        {'id': 'd2', 'project_id': 'p', 'deposit_date': '2026-08-01', 'amount': 2100, 'transaction_type': 'deposit'},
        {'id': 'd3', 'project_id': 'p', 'deposit_date': '2026-08-03', 'amount': -1050, 'transaction_type': 'reversal'},
        {'id': 'd4', 'project_id': 'p', 'deposit_date': '2026-09-01', 'amount': 1050, 'transaction_type': 'deposit'},
    ]
    entries = [dict(id=f'e{i}', project_id='p', entry_date=day, line_amount=gross, item_hours='0.5', quantity=2,
                    person_name='虛構使用者', item_name='測試項目', coach_id='c', created_at='2026-10-01')
               for i, (day, gross) in enumerate([('2026-07-30', 1050), ('2026-08-01', 2100), ('2026-08-31', 3150), ('2026-09-01', 4200)])]
    return project, entries, deposits


class ReportTests(unittest.TestCase):
    def report(self, mode='range', **kwargs):
        p, e, d = fixtures()
        return execution_report(p, e, d, {'c': '測試教練'}, tax, mode=mode, end=kwargs.pop('end', '2026-08-31'),
                                start=kwargs.pop('start', '2026-08-01'), **kwargs)

    def test_range_opening_excludes_first_day_and_future(self):
        report = self.report()
        self.assertEqual(report['opening'], 9450)
        self.assertEqual(report['gross'], 5250)
        self.assertEqual(report['net'], 5000)
        self.assertNotIn('前期餘額', report['frame'].columns)

    def test_cutoff_is_cumulative_with_month_opening(self):
        report = self.report('cutoff', start=None)
        self.assertEqual(report['opening'], 9450)
        self.assertEqual(report['gross'], 6300)
        self.assertEqual(report['deposited_gross'], 11550)
        self.assertEqual(report['balance_gross'], 5250)
        self.assertEqual(report['balance_net'], 5000)
        self.assertEqual(report['frame']['總儲值金額(含稅)'].dropna().tolist(), [11550])
        self.assertEqual(len(report['frame']), 3)
        self.assertEqual(report['frame']['教練'].unique().tolist(), ['測試教練'])

    def test_next_month_opening_includes_signed_reversal_and_usage(self):
        self.assertEqual(self.report(end='2026-09-30', start='2026-09-01')['opening'], 5250)

    def test_history_uses_entry_date_not_created_date(self):
        self.assertEqual(self.report()['gross'], 5250)

    def test_hours_multiply_quantity_and_no_unit_string(self):
        self.assertEqual(self.report()['frame']['時數'].tolist(), [1, 1])

    def test_exact_columns_range(self):
        self.assertEqual(self.report()['frame'].columns.tolist(), ['執行日期', '專案名稱', '教練', '使用者', '執行項目', '時數', '金額（含稅）', '金額（未稅）'])

    def test_exact_columns_cutoff(self):
        self.assertEqual(self.report('cutoff')['frame'].columns.tolist(), ['專案名稱', '總儲值金額(含稅)', '日期', '教練', '使用者', '項目', '時數', '金額（含稅）', '金額（未稅）'])

    def test_range_project_name_on_every_row_without_affecting_totals(self):
        report = self.report()
        self.assertEqual(report['frame']['專案名稱'].tolist(), ['虛構儲值專案', '虛構儲值專案'])
        self.assertEqual((report['gross'], report['net'], report['opening']), (5250, 5000, 9450))
        empty = self.report(start='2026-08-15', end='2026-08-20')['frame']
        self.assertEqual(empty.columns.tolist(), report['frame'].columns.tolist())

    def test_cutoff_label_renamed_and_no_duplicate_tax_suffix(self):
        frame = project_report_frame(self.report('cutoff'), 'cutoff')
        self.assertIn('總儲值金額(含稅)', frame.columns)
        self.assertNotIn('期初/期間儲值金額', frame.columns)
        self.assertEqual(frame['總儲值金額(含稅)'].dropna().tolist(), [11550])
        formats = project_excel_number_formats()
        self.assertIn('總儲值金額(含稅)', formats)
        self.assertFalse(any('期初/期間儲值金額' in name for name in formats))

    def test_oldest_dates_first_and_typed(self):
        self.assertEqual(self.report()['frame']['執行日期'].tolist(), [date(2026, 8, 1), date(2026, 8, 31)])

    def test_midmonth_opening_and_no_double_count(self):
        report = self.report(start='2026-08-15')
        self.assertEqual(report['opening'], 8400)
        self.assertEqual(report['gross'], 3150)

    def test_empty_period_keeps_opening(self):
        report = self.report(start='2026-08-15', end='2026-08-20')
        self.assertTrue(report['frame'].empty)
        self.assertEqual(report['gross'], 0)
        self.assertEqual(report['opening'], 8400)

    def test_filter_excludes_another_project(self):
        p, e, d = fixtures()
        e.append({**e[1], 'id': 'other', 'project_id': 'other'})
        self.assertEqual(execution_report(p, e, d, {}, tax, mode='range', start='2026-08-01', end='2026-08-31')['gross'], 5250)

    def test_unfunded_columns_and_note(self):
        p, e, d = fixtures(); p['funding_type'] = 'unfunded'; e[1]['note'] = '測試備註'
        report = execution_report(p, e, [], {'c': '測試教練'}, tax, mode='unfunded', start='2026-08-01', end='2026-08-31')
        self.assertEqual(report['frame'].columns.tolist(), ['日期', '專案名稱', '教練', '使用者', '項目', '時數', '金額（含稅）', '備註'])
        self.assertIn('測試備註', report['frame']['備註'].tolist())
        self.assertIsNone(report['opening'])

    def test_deposits_signed_and_no_stored_amount_double_count(self):
        p, e, d = fixtures(); report = deposit_report(p, d, tax, '2026-08-31')
        self.assertEqual(report['gross'], 11550)
        self.assertEqual(report['net'], 11000)
        self.assertEqual(report['frame']['類型'].tolist(), ['期初儲值', '後續儲值', '沖銷'])
        self.assertEqual(report['frame'].columns.tolist(), ['日期', '專案', '類型', '金額（含稅）', '金額（未稅）', '備註'])

    def test_missing_or_negative_data_are_errors(self):
        for field, bad in [('line_amount', None), ('item_hours', None), ('quantity', 0), ('line_amount', -1), ('line_amount', 'NaN')]:
            p, e, d = fixtures(); e[1][field] = bad
            with self.subTest(field=field, bad=bad), self.assertRaises(ValueError):
                execution_report(p, e, d, {}, tax, mode='range', start='2026-08-01', end='2026-08-31')

    def test_invalid_period_and_funding_rejected(self):
        with self.assertRaises(ValueError): self.report(start='2026-09-01')
        with self.assertRaises(ValueError): self.report('unfunded')

    def test_reversal_positive_is_error(self):
        p, e, d = fixtures(); d[2]['amount'] = 1050
        with self.assertRaises(ValueError): deposit_report(p, d, tax, '2026-08-31')
        with self.assertRaises(ValueError):
            execution_report(p, e, d, {}, tax, mode='cutoff', end='2026-08-31')

    def test_october_opening_includes_september_last_day_only(self):
        p, e, d = fixtures()
        d.extend([dict(d[1], id='last', deposit_date='2026-09-30', amount=2100),
                  dict(d[1], id='first', deposit_date='2026-10-01', amount=10500)])
        e.extend([dict(e[0], id='last', entry_date='2026-09-30', line_amount=1050),
                  dict(e[0], id='first', entry_date='2026-10-01', line_amount=2100)])
        report = execution_report(p, e, d, {'c': '測試教練'}, tax, mode='range', start='2026-10-01', end='2026-10-31')
        self.assertEqual(report['opening'], 3150)
        self.assertEqual(report['gross'], 2100)
        self.assertEqual(report['frame']['教練'].tolist(), ['測試教練'])

    def test_cutoff_with_deposits_without_execution(self):
        p, e, d = fixtures()
        report = execution_report(p, [], d, {}, tax, mode='cutoff', end='2026-08-31')
        self.assertTrue(report['frame'].empty)
        self.assertEqual(report['balance_gross'], 11550)
        self.assertEqual(report['balance_net'], 11000)

    def test_fully_used_balance_does_not_create_rounding_residual(self):
        p, e, d = fixtures()
        d = [dict(d[0], amount=3380)]
        e = [dict(e[1], line_amount=1690), dict(e[2], line_amount=1690)]
        report = execution_report(p, e, d, {}, tax, mode='cutoff', end='2026-08-31')
        self.assertNotEqual(tax(3380, '未稅'), report['net'])
        self.assertEqual(report['balance_gross'], 0)
        self.assertEqual(report['balance_net'], 0)

    def test_negative_balance_is_not_silently_clamped(self):
        p, e, d = fixtures()
        report = execution_report(p, e, d, {}, tax, mode='cutoff', end='2026-09-30')
        self.assertEqual(report['balance_gross'], 2100)
        report = execution_report(p, e, [dict(d[0], amount=1050)], {}, tax, mode='cutoff', end='2026-08-31')
        self.assertEqual(report['balance_gross'], -5250)
        self.assertEqual(report['balance_net'], -5000)

    def test_coach_unknown_and_unassigned_are_explicit(self):
        p, e, d = fixtures(); e[1]['coach_id'] = None
        report = execution_report(p, e, d, {}, tax, mode='range', start='2026-08-01', end='2026-08-31')
        self.assertEqual(report['frame']['教練'].tolist(), ['未指定', '未知教練'])

    def test_raw_inputs_not_modified(self):
        p, e, d = fixtures(); before = deepcopy((p, e, d))
        execution_report(p, e, d, {}, tax, mode='range', start='2026-08-01', end='2026-08-31')
        self.assertEqual((p, e, d), before)

    def test_net_metric_exactly_matches_displayed_lines(self):
        p, e, d = fixtures(); e[1]['line_amount'] = '1690.25'; e[2]['line_amount'] = '1690.25'
        report = execution_report(p, e, d, {}, tax, mode='range', start='2026-08-01', end='2026-08-31')
        self.assertEqual(report['net'], int(report['frame']['金額（未稅）'].sum()))
        self.assertEqual(report['gross'], 3380.5)

    def test_excel_dates_totals_and_summary(self):
        from openpyxl import load_workbook
        tree = ast.parse((ROOT / 'app.py').read_text(encoding='utf-8'))
        excel = next(x for x in tree.body if isinstance(x, ast.FunctionDef) and x.name == '_excel_bytes')
        ns = {'BytesIO': BytesIO, 'pd': pd}
        exec(compile(ast.Module(body=[excel], type_ignores=[]), '<excel>', 'exec'), ns)
        report = self.report()
        payload = ns['_excel_bytes']({'查詢結果': report['frame'], '查詢摘要': pd.DataFrame([{'專案執行總計（未稅）': report['net'], '專案執行總計（含稅）': report['gross']}])}, number_formats=project_excel_number_formats())
        wb = load_workbook(BytesIO(payload), data_only=True)
        ws = wb['查詢結果']
        self.assertIsInstance(ws['A2'].value, datetime)
        self.assertEqual(ws['A2'].number_format, 'yyyy-mm-dd')
        self.assertEqual(ws.freeze_panes, 'A2')
        self.assertEqual([cell.value for cell in ws[1]], report['frame'].columns.tolist())
        self.assertEqual(ws['B2'].value, '虛構儲值專案')
        self.assertEqual(ws['B3'].value, '虛構儲值專案')
        self.assertEqual(ws['G2'].number_format, project_excel_number_formats()['金額（含稅）'])
        self.assertEqual(ws['H2'].number_format, project_excel_number_formats()['金額（未稅）'])
        self.assertEqual(ws['F2'].number_format, '0.##')
        self.assertNotIn('.00', wb['查詢摘要']['B2'].number_format)
        self.assertEqual(wb['查詢摘要']['A1'].value, '專案執行總計（未稅）')
        self.assertEqual(wb['查詢摘要']['B1'].value, '專案執行總計（含稅）')
        self.assertEqual(sum(ws.cell(i, 8).value for i in (2, 3)), wb['查詢摘要']['A2'].value)
        self.assertNotIn('前期餘額', [cell.value for cell in ws[1]])

    def test_excel_whole_display_preserves_fractional_source(self):
        from openpyxl import load_workbook
        tree = ast.parse((ROOT / 'app.py').read_text(encoding='utf-8'))
        excel = next(x for x in tree.body if isinstance(x, ast.FunctionDef) and x.name == '_excel_bytes')
        ns = {'BytesIO': BytesIO, 'pd': pd}
        exec(compile(ast.Module(body=[excel], type_ignores=[]), '<excel>', 'exec'), ns)
        frame = pd.DataFrame([{'專案餘額總計（含稅）': 3380.5, '總儲值金額(含稅)': 3380.5}])
        wb = load_workbook(BytesIO(ns['_excel_bytes']({'查詢摘要': frame}, number_formats=project_excel_number_formats())), data_only=True)
        for cell in wb['查詢摘要'][2]:
            self.assertEqual(cell.value, 3380.5)
            self.assertNotIn('.00', cell.number_format)

    def test_footer_is_last_without_changing_detail_or_totals(self):
        report = self.report('cutoff', start=None)
        before = report['frame'].copy(deep=True)
        frame = project_report_frame(report, 'cutoff')
        pd.testing.assert_frame_equal(report['frame'], before)
        self.assertEqual(len(frame), 4)
        self.assertEqual(frame.iloc[-1]['專案名稱'], '專案餘額總計')
        self.assertEqual(frame.iloc[-1]['金額（含稅）'], 5250)
        self.assertEqual(frame.iloc[-1]['金額（未稅）'], 5000)
        self.assertTrue(pd.isna(frame.iloc[-1]['日期']))
        self.assertTrue(pd.isna(frame.iloc[-1]['時數']))
        self.assertEqual(frame.iloc[:-1]['日期'].tolist(), [date(2026, 7, 30), date(2026, 8, 1), date(2026, 8, 31)])
        self.assertEqual(report['gross'], 6300)
        self.assertEqual(report['net'], 6000)
        self.assertEqual(len(report['frame']), 3)
        pd.testing.assert_frame_equal(project_report_frame(report, 'cutoff'), frame)

    def test_empty_cutoff_has_only_balance_row(self):
        p, e, d = fixtures()
        report = execution_report(p, [], d, {}, tax, mode='cutoff', end='2026-08-31')
        frame = project_report_frame(report, 'cutoff')
        self.assertEqual(len(frame), 1)
        self.assertEqual(frame.iloc[0]['專案名稱'], '專案餘額總計')
        self.assertEqual(frame.iloc[0]['金額（含稅）'], 11550)
        self.assertEqual(frame.iloc[0]['金額（未稅）'], 11000)
        self.assertTrue(report['frame'].empty)
        self.assertEqual(report['gross'], 0)

    def test_noncutoff_presentation_has_no_balance_footer(self):
        for mode in ['range', 'unfunded', 'deposit']:
            if mode == 'deposit':
                p, e, d = fixtures(); report = deposit_report(p, d, tax, '2026-08-31')
            elif mode == 'unfunded':
                p, e, d = fixtures(); p['funding_type'] = 'unfunded'
                report = execution_report(p, e, [], {}, tax, mode=mode, start='2026-08-01', end='2026-08-31')
            else:
                report = self.report()
            pd.testing.assert_frame_equal(project_report_frame(report, mode), report['frame'])
            dates = report['frame']['執行日期' if mode == 'range' else '日期'].tolist()
            self.assertEqual(dates, sorted(dates))

    def test_all_currency_formats_are_whole_and_hours_keep_fraction(self):
        formats = project_excel_number_formats()
        self.assertEqual(formats.pop('時數'), '0.##')
        self.assertTrue(all('.00' not in fmt for fmt in formats.values()))

    def test_same_day_order_is_stable_and_oldest_created_first(self):
        p, e, d = fixtures()
        e = [dict(e[1], id='b', person_name='第二筆', created_at='2026-08-01T10:00:00'),
             dict(e[1], id='a', person_name='第一筆', created_at='2026-08-01T09:00:00')]
        report = execution_report(p, e, d, {}, tax, mode='range', start='2026-08-01', end='2026-08-31')
        self.assertEqual(report['frame']['使用者'].tolist(), ['第一筆', '第二筆'])

    def test_excel_footer_dates_precision_and_detail_count(self):
        from openpyxl import load_workbook
        tree = ast.parse((ROOT / 'app.py').read_text(encoding='utf-8'))
        excel = next(x for x in tree.body if isinstance(x, ast.FunctionDef) and x.name == '_excel_bytes')
        ns = {'BytesIO': BytesIO, 'pd': pd}
        exec(compile(ast.Module(body=[excel], type_ignores=[]), '<excel>', 'exec'), ns)
        report = self.report('cutoff', start=None)
        frame = project_report_frame(report, 'cutoff')
        payload = ns['_excel_bytes']({'查詢結果': frame, '查詢摘要': pd.DataFrame([{'資料筆數': len(report['frame']), '專案執行金額合計（含稅）': report['gross']}])},
                                     number_formats=project_excel_number_formats(), fit_display_width=True)
        wb = load_workbook(BytesIO(payload), data_only=True)
        ws = wb['查詢結果']
        self.assertEqual(ws.max_row, 5)
        self.assertEqual(ws['B1'].value, '總儲值金額(含稅)')
        self.assertEqual(ws['B2'].value, 11550)
        self.assertEqual(ws['B2'].number_format, project_excel_number_formats()['總儲值金額(含稅)'])
        self.assertEqual(ws['A5'].value, '專案餘額總計')
        self.assertEqual((ws['H5'].value, ws['I5'].value), (5250, 5000))
        self.assertIsNone(ws['C5'].value)
        self.assertIsInstance(ws['C2'].value, datetime)
        self.assertLess(ws['C2'].value, ws['C4'].value)
        for row in range(2, 6):
            for col in (8, 9):
                self.assertEqual(ws.cell(row, col).data_type, 'n')
                self.assertNotIn('.00', ws.cell(row, col).number_format)
        self.assertEqual(wb['查詢摘要']['A2'].value, 3)
        self.assertEqual(wb['查詢摘要']['B2'].value, 6300)
        self.assertGreaterEqual(ws.column_dimensions['A'].width, 14)

    def test_project_export_width_only_affects_opted_in_exports(self):
        from openpyxl import load_workbook
        tree = ast.parse((ROOT / 'app.py').read_text(encoding='utf-8'))
        excel = next(x for x in tree.body if isinstance(x, ast.FunctionDef) and x.name == '_excel_bytes')
        ns = {'BytesIO': BytesIO, 'pd': pd}
        exec(compile(ast.Module(body=[excel], type_ignores=[]), '<excel>', 'exec'), ns)
        data = {'結果': pd.DataFrame([{'專案名稱': '專案餘額總計', '金額（含稅）': 3380.5}])}
        original = load_workbook(BytesIO(ns['_excel_bytes'](data)), data_only=True)
        project = load_workbook(BytesIO(ns['_excel_bytes'](data, number_formats=project_excel_number_formats(), fit_display_width=True)), data_only=True)
        self.assertGreater(project['結果'].column_dimensions['A'].width, original['結果'].column_dimensions['A'].width)
        self.assertEqual(project['結果']['B2'].value, original['結果']['B2'].value)
        self.assertEqual(project['結果']['B2'].value, 3380.5)
        self.assertNotIn('.00', project['結果']['B2'].number_format)
        source = (ROOT / 'app.py').read_text(encoding='utf-8')
        self.assertIn('number_formats=project_excel_number_formats(),fit_display_width=True', source)

    def test_parent_admin_gate_and_dynamic_project_tab(self):
        source = (ROOT / 'app.py').read_text(encoding='utf-8')
        page = source[source.index('def general_queries_page'):source.index('def financial_report_page')]
        self.assertLess(page.index('if me["role"]!="admin"'), page.index('render_project_queries(me'))
        self.assertIn('if project_tab.open:', page)
        self.assertIn('on_change="rerun"', page)


if __name__ == '__main__':
    unittest.main()
