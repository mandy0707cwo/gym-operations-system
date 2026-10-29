"""只驗證顯示順序；虛構資料不連正式系統。"""
import unittest
from io import BytesIO
import pandas as pd
from pandas.testing import assert_frame_equal
from openpyxl import load_workbook
from monthly_report_layout import sort_account_rows
import test_finance_report_columns_v38 as v38
from test_financial_backup import load_functions


class AccountSortTests(unittest.TestCase):
    def test_natural_order_blank_last_stable_and_no_values_changed(self):
        frame=pd.DataFrame({'帳號':['smc010','SMC002','smc001',None,'smc2','', 'smc000'],
                            '金額':[10,20,-3,4,5,6,7],'時數':[0.5,2,1,0,3,4,5]})
        before=frame.copy(deep=True);result=sort_account_rows(frame)
        # Pandas可能將輸入的None正規化為NaN；比對原列可同時證明缺值與數值未改。
        assert_frame_equal(result,frame.iloc[[6,2,1,4,0,3,5]].reset_index(drop=True))
        self.assertEqual(result['帳號'].fillna('').tolist(),['smc000','smc001','SMC002','smc2','smc010','',''])
        self.assertEqual(result['金額'].tolist(),[7,-3,20,5,10,4,6])
        self.assertEqual(result['金額'].sum(),frame['金額'].sum())
        self.assertEqual(result['時數'].sum(),frame['時數'].sum())
        assert_frame_equal(frame,before)

    def test_month_first_then_account_and_same_account_details_stable(self):
        frame=pd.DataFrame({'報表月份':['2026-10','2026-09','2026-10','2026-10'],
                            '帳號':['smc002','smc010','smc001','smc001'],'筆次':[0,1,2,3]})
        self.assertEqual(sort_account_rows(frame)['筆次'].tolist(),[1,2,3,0])

    def test_empty_and_frames_without_account_are_unchanged(self):
        for frame in (pd.DataFrame(columns=['帳號','教練']),pd.DataFrame({'教練':['乙','甲'],'金額':[2,1]})):
            assert_frame_equal(sort_account_rows(frame),frame)

    def test_monthly_screen_excel_and_backup_all_account_tables_sorted(self):
        fixture=v38.FinanceReportColumnsTests();data=fixture.bonus_data()
        data['profiles'][0]['username']='smc010';data['profiles'][1]['username']='SMC002'
        app=fixture.monthly(data);frames=app.session_state['exports']
        backup=load_functions()['_financial_backup_frames'](data)
        for key in ('每月教練時數','每月教練營收','談單獎金總計','結單獎金總計'):
            self.assertEqual(frames[key]['帳號'].tolist(),['SMC002','smc010'],key)
        for key in ('財務-教練時數','財務-教練營收','財務-談單獎金總計','財務-結單獎金總計'):
            self.assertEqual(backup[key]['帳號'].tolist(),['SMC002','smc010'],key)
        # 畫面與下載使用同一份已排序資料；不能只改瀏覽器欄位排序。
        for frame in frames.values():
            if '帳號' in frame.columns and not frame.empty:
                self.assertTrue(any(frame.equals(item.value) for item in app.dataframe))
        book=load_workbook(BytesIO(app.session_state['export_bytes']),data_only=True)
        self.assertEqual([book['每月教練營收'].cell(i,1).value for i in (2,3)],['SMC002','smc010'])
        by_account=frames['談單獎金總計'].set_index('帳號')
        self.assertEqual(by_account.loc['smc010','可計成交金額'],2000)
        self.assertEqual(by_account.loc['smc010','不計成交金額'],1000)
        self.assertEqual(by_account.loc['smc010','獎金合計'],200)

    def test_pending_details_sorted_and_same_account_date_order_preserved(self):
        fixture=v38.FinanceReportColumnsTests();data=fixture.bonus_data()
        data['profiles'][0]['username']='smc010';data['profiles'][1]['username']='SMC002'
        data['bonus_rules'][0]['effective_from']='2026-10-01'
        app=fixture.monthly(data);frames=app.session_state['exports']
        backup=load_functions()['_financial_backup_frames'](data)
        for key in ('談單待確認','結單待確認'):
            self.assertEqual(frames[key]['帳號'].tolist(),['SMC002','smc010','smc010'])
            self.assertEqual(frames[key]['待確認成交金額（未稅）'].tolist(),[3000,2000,1000])
            assert_frame_equal(frames[key],backup['財務-'+key])


if __name__=='__main__':unittest.main()
