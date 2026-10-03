"""驗證實際付清尾差、截止日，以及真實查詢函式的畫面／下載一致性。"""
import ast
from copy import deepcopy
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest

import pandas as pd
from openpyxl import load_workbook

APP_PATH = Path(__file__).resolve().parents[1] / "app.py"


def sample_purchase():
    return {"id": "p1", "total_amount": 15840, "total_sessions": 12, "member_id": "m1", "coach_id": "c1",
            "course_name": "運動恢復", "purchase_date": "2026-04-14", "expiry_date": "2027-04-14",
            "payment_plan": "installment", "installment_count": 3}


def sample_payments(amounts=(5280, 5280, 5280)):
    dates = ["2026-04-14", "2026-05-08", "2026-07-03"]
    return [{"id": f"r{i}", "purchase_id": "p1", "amount": amount, "paid_date": dates[i],
             "created_at": f"2026-08-13T09:0{i}:00", "installment_no": i + 1}
            for i, amount in enumerate(amounts)]


def load_helpers():
    tree = ast.parse(APP_PATH.read_text(encoding="utf-8"))
    wanted = {"_tax_display_amount", "_allocate_payment_net_amounts", "usage_net_amount", "_excel_bytes"}
    selected = [x for x in tree.body if isinstance(x, ast.FunctionDef) and x.name in wanted]
    ns = {"Decimal": Decimal, "ROUND_HALF_UP": ROUND_HALF_UP, "pd": pd, "date": date, "BytesIO": BytesIO}
    exec(compile(ast.Module(body=selected, type_ignores=[]), str(APP_PATH), "exec"), ns)
    return ns


class AllocationTests(unittest.TestCase):
    def setUp(self):
        self.allocate = load_helpers()["_allocate_payment_net_amounts"]

    def test_verified_three_installment_example(self):
        rows = self.allocate([sample_purchase()], sample_payments())
        self.assertEqual([x["_net_amount"] for x in rows], [5029, 5029, 5028])
        self.assertEqual(sum(x["_net_amount"] for x in rows), 15086)
        self.assertEqual(sum(x["amount"] for x in rows), 15840)

    def test_prior_period_is_unchanged_before_payoff(self):
        self.assertEqual([x["_net_amount"] for x in self.allocate([sample_purchase()], sample_payments()[:2])], [5029, 5029])

    def test_last_installment_number_does_not_hide_underpayment(self):
        rows = self.allocate([sample_purchase()], sample_payments((5280, 5280, 4000)))
        self.assertEqual(rows[-1]["_net_amount"], 3810)
        self.assertEqual(sum(x["amount"] for x in rows), 14560)
        self.assertNotEqual(sum(x["_net_amount"] for x in rows), 15086)

    def test_supplement_absorbs_tail_when_actually_paid(self):
        payments = sample_payments((5280, 5280, 4000))
        payments.append({"id": "r4", "purchase_id": "p1", "amount": 1280, "paid_date": "2026-07-04", "payment_kind": "supplement"})
        rows = self.allocate([sample_purchase()], payments)
        self.assertEqual(rows[-1]["_net_amount"], 1218)
        self.assertEqual(sum(x["_net_amount"] for x in rows), 15086)

    def test_overpayment_not_capped_or_hidden(self):
        rows = self.allocate([sample_purchase()], sample_payments((5280, 5280, 5385)))
        self.assertEqual(rows[-1]["_net_amount"], 5128)
        self.assertEqual(sum(x["_net_amount"] for x in rows), 15186)
        self.assertEqual(sum(x["amount"] for x in rows), 15945)

    def test_post_payoff_extra_receipt_preserved(self):
        payments = sample_payments() + [{"id": "r4", "purchase_id": "p1", "amount": 105, "paid_date": "2026-07-04"}]
        self.assertEqual(self.allocate([sample_purchase()], payments)[-1]["_net_amount"], 100)

    def test_input_not_mutated_and_order_deterministic(self):
        payments = sample_payments()
        original = deepcopy(payments)
        self.assertEqual(self.allocate([sample_purchase()], payments[::-1]), self.allocate([sample_purchase()], payments))
        self.assertEqual(payments, original)

    def test_same_day_tie_uses_created_at_then_id(self):
        payments = sample_payments()
        for row in payments:
            row["paid_date"] = "2026-04-14"
            row["created_at"] = "2026-04-14T00:00:00"
        rows = self.allocate([sample_purchase()], payments[::-1])
        self.assertEqual([x["id"] for x in rows], ["r0", "r1", "r2"])
        self.assertEqual(rows[-1]["_net_amount"], 5028)

    def test_purchases_independent_and_full_payment(self):
        p2 = {**sample_purchase(), "id": "p2", "total_amount": 2100}
        rows = self.allocate([sample_purchase(), p2], sample_payments() + [{"id": "other", "purchase_id": "p2", "amount": 2100}])
        self.assertEqual(next(x["_net_amount"] for x in rows if x["id"] == "other"), 2000)

    def test_anomalies_raise_instead_of_silently_zeroing(self):
        with self.assertRaises(ValueError): self.allocate([], sample_payments())
        with self.assertRaises(ValueError): self.allocate([sample_purchase()], sample_payments((-1, 0, 0)))
        p = {**sample_purchase(), "total_amount": 10}
        many = [{"id": str(i), "purchase_id": "p1", "amount": 1, "paid_date": f"2026-04-{i+1:02d}"} for i in range(10)]
        many.append({"id": "zero", "purchase_id": "p1", "amount": 0, "paid_date": "2026-04-11"})
        rows = self.allocate([p], many)
        self.assertEqual(sum(x["_net_amount"] for x in rows), 10)
        self.assertEqual(rows[-1]["_net_amount"], 0)
        fractional = [{"id": str(i), "purchase_id": "p1", "amount": "0.60", "paid_date": "2026-04-14", "created_at": f"{i:03d}"} for i in range(50)]
        with self.assertRaises(ValueError):
            self.allocate([{**sample_purchase(), "total_amount": "30.00"}], fractional)

    def test_no_payment_returns_empty(self):
        self.assertEqual(self.allocate([sample_purchase()], []), [])


class Query:
    def __init__(self, records): self.records = list(records)
    def select(self, *_): return self
    def order(self, *_, **__): return self
    def in_(self, field, values):
        self.records = [x for x in self.records if x.get(field) in values]; return self
    def lte(self, field, value):
        self.records = [x for x in self.records if str(x.get(field) or "") <= value]; return self
    def gte(self, field, value):
        self.records = [x for x in self.records if str(x.get(field) or "") >= value]; return self


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.ns = load_helpers()
        self.tables = {"purchases": [sample_purchase()], "purchase_payments": sample_payments(), "course_terminations": [],
                       "session_usages": [{"id": f"u{i}", "purchase_id": "p1", "usage_date": "2026-07-15",
                                          "deducted_amount": 1320, "deducted_net_amount": 1258 if i in (3, 10) else 1257} for i in range(12)]}
        self.frame = None; self.download = None; self.metrics = {}
        metric = lambda label, value, **_: self.metrics.update({label: value})
        st = SimpleNamespace(columns=lambda _: [SimpleNamespace(metric=metric)] * 2, metric=metric,
            caption=lambda *_: None, info=lambda *_: None, dataframe=lambda frame, **_: setattr(self, "frame", frame),
            column_config=SimpleNamespace(DateColumn=lambda **_: None, NumberColumn=lambda **_: None))
        self.ns.update({"client": lambda: SimpleNamespace(table=lambda table: Query(self.tables[table])),
            "paged_rows": lambda builder: builder().records, "member_maps": lambda: {"m1": "測試會員"},
            "purchase_codes": lambda: {"p1": "20260414-003"}, "coach_name_map": {"c1": "測試教練"},
            "coaches": {"測試教練": "c1"}, "st": st,
            "download_frame": lambda label, frame, filename: setattr(self, "download", frame.copy())})
        tree = ast.parse(APP_PATH.read_text(encoding="utf-8"))
        page = next(x for x in tree.body if isinstance(x, ast.FunctionDef) and x.name == "general_queries_page")
        renderers = [x for x in page.body if isinstance(x, ast.FunctionDef) and x.name in {"render_prepaid_results", "build_purchase_results", "render_purchase_results", "render_balance_results"}]
        exec(compile(ast.Module(body=renderers, type_ignores=[]), str(APP_PATH), "exec"), self.ns)

    def test_last_period_receipt_and_excel_have_correct_tail(self):
        self.ns["render_prepaid_results"](start=date(2026, 7, 1), end=date(2026, 7, 31), key="july")
        self.assertEqual(self.frame["預收金額（未稅）"].tolist(), [5028])
        self.assertEqual(self.frame["分期狀態"].tolist(), ["付清"])
        self.assertEqual(self.metrics["預收金額總計（未稅）"], "$ 5,028")
        pd.testing.assert_frame_equal(self.frame, self.download)
        workbook = load_workbook(BytesIO(self.ns["_excel_bytes"]({"實際預收收入": self.download})), data_only=True)
        self.assertEqual(workbook.active["H2"].value, 5028)
        self.assertEqual(workbook.active["A2"].data_type, "d")

    def test_early_cutoff_excludes_future_payoff(self):
        self.ns["render_prepaid_results"](start=date(2026, 5, 1), end=date(2026, 5, 31), key="may")
        self.assertEqual(self.frame["預收金額（未稅）"].tolist(), [5029])
        self.assertEqual(self.frame["分期狀態"].tolist(), ["未付清"])

    def test_purchase_and_balance_net_totals_match(self):
        frame = self.ns["build_purchase_results"](self.tables["purchases"], date(2026, 10, 3))
        self.assertEqual(frame.iloc[0]["已收金額（未稅）"], 15086)
        self.ns["render_balance_results"](date(2026, 10, 3), "", "未稅", "net")
        self.assertEqual(self.frame.iloc[0]["已收金額（未稅）"], 15086)
        self.assertEqual(self.frame.iloc[0]["銷課金額（未稅）"], 15086)
        self.assertEqual(self.frame.iloc[0]["餘額（未稅）"], 0)
        pd.testing.assert_frame_equal(self.frame, self.download)

    def test_gross_balance_unchanged(self):
        self.ns["render_balance_results"](date(2026, 10, 3), "", "含稅", "gross")
        self.assertEqual(self.frame.iloc[0]["已收金額（含稅）"], 15840)
        self.assertEqual(self.frame.iloc[0]["餘額（含稅）"], 0)

    def test_unpaid_filter_retains_real_shortfall(self):
        self.tables["purchase_payments"][-1]["amount"] = 4000
        frame = self.ns["build_purchase_results"](self.tables["purchases"], date(2026, 10, 3), payment_filter="未付清")
        self.assertEqual(len(frame), 1)
        self.assertEqual(frame.iloc[0]["已收金額（含稅）"], 14560)
        self.ns["render_balance_results"](date(2026, 10, 3), "", "未稅", "negative")
        self.assertLess(self.frame.iloc[0]["餘額（未稅）"], 0)

    def test_zero_results_have_zero_totals(self):
        self.ns["render_prepaid_results"](member_keyword="不存在", key="empty")
        self.assertIsNone(self.download)
        self.assertEqual(self.metrics["預收金額總計（未稅）"], "$ 0")

    def test_purchase_last_three_columns_and_excel_match(self):
        purchase = self.tables["purchases"][0]
        purchase.update(purchase_kind="first", referral="測試醫師", note="購課時備註\n第二行")
        frame = self.ns["build_purchase_results"](self.tables["purchases"], date(2026, 10, 3))
        self.assertEqual(frame.columns.tolist()[-3:], ["購買類型", "醫生轉介", "備註"])
        self.assertEqual(frame.iloc[0]["購買類型"], "首次購買")
        self.assertEqual(frame.iloc[0]["醫生轉介"], "測試醫師")
        self.assertEqual(frame.iloc[0]["備註"], "購課時備註\n第二行")
        self.ns["render_purchase_results"](frame, "columns")
        pd.testing.assert_frame_equal(self.frame, self.download)
        self.assertEqual(self.metrics["成交金額總計（未稅）"], "$ 15,086")
        self.assertEqual(self.metrics["成交金額總計（含稅）"], "$ 15,840")
        sheet = load_workbook(BytesIO(self.ns["_excel_bytes"]({"成交總表": self.download})), data_only=True).active
        self.assertEqual([cell.value for cell in sheet[1]][-3:], ["購買類型", "醫生轉介", "備註"])
        self.assertEqual([cell.value for cell in sheet[2]][-3:], ["首次購買", "測試醫師", "購課時備註\n第二行"])
        self.assertEqual(sheet["A2"].data_type, "d")

    def test_purchase_types_and_nulls_are_preserved(self):
        purchase = self.tables["purchases"][0]
        for kind, label in (("first", "首次購買"), ("renewal", "續課"), ("legacy", "legacy"), (None, "")):
            with self.subTest(kind=kind):
                purchase.update(purchase_kind=kind, referral=None, note=None)
                row = self.ns["build_purchase_results"](self.tables["purchases"], date(2026, 10, 3)).iloc[0]
                self.assertEqual(row["購買類型"], label)
                self.assertEqual(row["醫生轉介"], "")
                self.assertEqual(row["備註"], "")
                self.assertEqual(row["已收金額（未稅）"], 15086)

    def test_empty_purchase_results_keep_columns(self):
        frame = self.ns["build_purchase_results"]([], date(2026, 10, 3))
        self.assertTrue(frame.empty)
        self.assertEqual(frame.columns.tolist()[-3:], ["購買類型", "醫生轉介", "備註"])
        self.ns["render_purchase_results"](frame, "empty")
        self.assertIsNone(self.download)
        self.assertEqual(self.metrics["成交金額總計（未稅）"], "$ 0")

    def test_both_purchase_search_paths_select_new_fields(self):
        tree = ast.parse(APP_PATH.read_text(encoding="utf-8"))
        page = next(x for x in tree.body if isinstance(x, ast.FunctionDef) and x.name == "general_queries_page")
        expected = "id,member_id,coach_id,course_name,total_sessions,total_amount,purchase_date,expiry_date,payment_plan,installment_count,purchase_kind,referral,note,created_at"
        selects = [x for x in ast.walk(page) if isinstance(x, ast.Call) and isinstance(x.func, ast.Attribute)
                   and x.func.attr == "select" and x.args and isinstance(x.args[0], ast.Constant) and x.args[0].value == expected]
        self.assertEqual(len(selects), 2)


if __name__ == "__main__": unittest.main()
