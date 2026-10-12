"""中止記錄與正式快照驗證，僅使用虛構資料，不連線正式資料庫。"""
import ast
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
import copy
import hashlib
import unittest
import pandas as pd
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
COLUMNS=["中止日期","購買_id","會員","課程","中止類型","剩餘堂數","餘額","逾期入帳","退費手續費","實際退費","結單計算","中止原因"]
MONEY=["餘額","逾期入帳","退費手續費","實際退費"]

def functions(path):
    tree=ast.parse(path.read_text(encoding="utf-8"))
    selected=[x for x in tree.body if isinstance(x,ast.FunctionDef) and x.name in
              {"_course_termination_history_frame","_tax_display_amount","_excel_bytes"}]
    ns={"pd":pd,"Decimal":Decimal,"ROUND_HALF_UP":ROUND_HALF_UP,"BytesIO":BytesIO}
    exec(compile(ast.Module(body=selected,type_ignores=[]),str(path),"exec"),ns)
    return ns

class TerminationRecordsTests(unittest.TestCase):
    def setUp(self):
        self.ns=functions(ROOT/"app.py")
        self.items=[{"purchase_id":"p","termination_date":"2026-10-02","termination_type":"refund",
                     "remaining_sessions":4,"remaining_amount":1050,"recognized_amount":210,
                     "fee_amount":210,"refund_amount":840,"completion_bonus_eligible":False,"reason":"虛構原因"}]
        self.purchases={"p":{"member_id":"m","course_name":"虛構課程"}}
        self.members={"m":"虛構會員"}
        self.codes={"p":"20260101-001"}

    def frame(self,mode):
        return self.ns["_course_termination_history_frame"](self.items,self.purchases,self.members,self.codes,mode)

    def test_gross_exact_columns_and_no_input_mutation(self):
        original=copy.deepcopy(self.items)
        frame=self.frame("含稅")
        self.assertEqual(list(frame.columns),COLUMNS)
        self.assertEqual(list(frame.iloc[0][MONEY]),[1050,0,210,840])
        self.assertEqual(frame.iloc[0]["購買_id"],"20260101-001")
        self.assertEqual(self.items,original)

    def test_net_uses_existing_rounding_and_refund_not_expired_income(self):
        frame=self.frame("未稅")
        self.assertEqual(list(frame.iloc[0][MONEY]),[1000,0,200,800])
        self.assertEqual(frame.iloc[0]["結單計算"],"不計算")
        self.items[0]["remaining_amount"]=1050.51
        self.assertEqual(self.frame("未稅").iloc[0]["餘額"],self.ns["_tax_display_amount"](1050.51,"未稅"))
        self.assertEqual(self.frame("含稅").iloc[0]["餘額"],1050.51)

    def test_expired_and_empty_records(self):
        self.items[0].update(termination_type="expired",recognized_amount=1050,fee_amount=0,
                            refund_amount=0,completion_bonus_eligible=True)
        self.assertEqual(list(self.frame("含稅").iloc[0][MONEY]),[1050,1050,0,0])
        self.assertEqual(list(self.frame("未稅").iloc[0][MONEY]),[1000,1000,0,0])
        self.assertEqual(self.frame("未稅").iloc[0]["結單計算"],"計算")
        self.items=[]
        self.assertEqual(list(self.frame("未稅").columns),COLUMNS)
        self.assertTrue(self.frame("未稅").empty)

    def test_excel_numeric_date_tax_tabs_and_format(self):
        frames={f"中止記錄({mode})":self.frame(mode) for mode in ["含稅","未稅"]}
        workbook=load_workbook(BytesIO(self.ns["_excel_bytes"](frames,number_formats={x:'#,##0;(#,##0)' for x in MONEY})))
        for mode,expected in [("含稅",[1050,0,210,840]),("未稅",[1000,0,200,800])]:
            sheet=workbook[f"中止記錄({mode})"]
            self.assertEqual([c.value for c in sheet[1]],COLUMNS)
            self.assertEqual([sheet.cell(2,COLUMNS.index(c)+1).value for c in MONEY],expected)
            self.assertIsInstance(sheet["A2"].value,datetime)
            self.assertEqual(sheet["A2"].number_format,"yyyy-mm-dd")
            self.assertEqual(sheet.freeze_panes,"A2")
            self.assertTrue(sheet.auto_filter.ref)
            for column in MONEY:
                cell=sheet.cell(2,COLUMNS.index(column)+1)
                self.assertEqual(cell.data_type,"n")
                self.assertEqual(cell.number_format,'#,##0;(#,##0)')

    def test_financial_entry_removed_but_monthly_and_backup_unchanged(self):
        source=(ROOT/"app.py").read_text(encoding="utf-8")
        tree=ast.parse(source)
        finance=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=="financial_report_page")
        self.assertNotIn("course_termination_report_page",ast.unparse(finance))
        menu=next(n for n in ast.walk(finance) if isinstance(n,ast.Call)
                  and isinstance(n.func,ast.Attribute) and n.func.attr=="segmented_control"
                  and n.args and isinstance(n.args[0],ast.Constant) and n.args[0].value=="報表分類")
        self.assertEqual(ast.literal_eval(menu.args[1]),["會員報表","其他報表","每月報表","教練查詢"])
        self.assertIn('combined_tabs=st.tabs(["每月銷課","每月專案銷課","每月課程中止"])',source)
        self.assertIn('"財務-課程中止":pd.DataFrame(termination_rows)',source)
        new_backup=next(n for n in tree.body if isinstance(n,ast.FunctionDef) and n.name=="_financial_backup_frames")
        self.assertEqual(hashlib.sha256(ast.dump(new_backup).encode()).hexdigest(),"03e6850180c90a28a3b75f43efacad03875dc3d2f5a8bd7a918fc01234b4b5ed")

    def test_formal_snapshot_preserves_paused_pilot_and_versions(self):
        formal=ROOT/"output/release_v1.13.9/app.py"
        # 正式 GitHub 根目錄本身就是已排除原型的程式，不要求下載本機工作快照。
        if not formal.exists():
            formal=ROOT/"app.py"
        tree=ast.parse((ROOT/"app.py").read_text(encoding="utf-8"))
        official=ast.parse(formal.read_text(encoding="utf-8"))
        actual={n.name:ast.dump(n) for n in tree.body if isinstance(n,ast.FunctionDef)}
        staged={n.name:ast.dump(n) for n in official.body if isinstance(n,ast.FunctionDef)}
        for name in staged:
            if name!="daily_page":
                self.assertEqual(actual[name],staged[name],name)
        self.assertEqual(hashlib.sha256(staged["daily_page"].encode()).hexdigest(),"d93997d9b7318aa06c430818d60202289b4f7e3dbbafd9fce2fa6bc63f19ff0c")
        self.assertEqual((ROOT/"VERSION").read_text().strip(),"v1.13.9")
        self.assertIn('secret("APP_VERSION") or "v1.13.9"',formal.read_text(encoding="utf-8"))

if __name__=="__main__":
    unittest.main()
